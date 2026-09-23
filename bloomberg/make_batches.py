"""
Erzeugt Teil-Vorlagen fuer den Bloomberg-Pull, falls das Tageslimit
("#N/A Daily Capacity") keinen Abzug aller 60 Ticker auf einmal erlaubt.

    py bloomberg/make_batches.py                    # 4 Bloecke, neueste Jahrgaenge zuerst
    py bloomberg/make_batches.py --vintages-per-batch 2
    py bloomberg/make_batches.py --out-dir bloomberg/batches

Aufteilung erfolgt JAHRGANGSWEISE: ein Jahrgang braucht alle vier Beine
(RB/HO x Jun/Dez), sonst faellt er in der Analyse komplett aus. Reihenfolge
absteigend, damit der erste Block schon den laufenden Jahrgang liefert.

Jede Datei ist eine vollwertige Vorlage (Blaetter Daten/Tickerliste/Anleitung),
enthaelt aber nur die Ticker ihres Blocks. Ablauf je Datei auf dem Bloomberg-PC:
oeffnen -> EINMAL Strg+Alt+F9 -> Strg+A, Strg+C, Inhalte einfuegen -> Werte -> speichern.
Danach die Blaetter "Daten" aller Bloecke als getrennte Blaetter in EINE Mappe
bloomberg_export.xlsx kopieren (der Loader liest alle Blaetter und mischt sie).
"""
from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "src"))
from spread import config  # noqa: E402

import make_template as mt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402


def batches(vintages_per_batch: int) -> list[list[int]]:
    """Jahrgaenge absteigend, in Bloecke zerlegt (neueste zuerst)."""
    vs = sorted(config.VINTAGES, reverse=True)
    return [vs[i:i + vintages_per_batch] for i in range(0, len(vs), vintages_per_batch)]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Teil-Vorlagen fuer den Bloomberg-Pull")
    ap.add_argument("--vintages-per-batch", type=int, default=4,
                    help="Jahrgaenge je Block (Standard 4 = 16 Ticker)")
    ap.add_argument("--out-dir", type=Path, default=HERE / "batches")
    ap.add_argument("--start", default=config.HISTORY_START.strftime("%Y%m%d"))
    ap.add_argument("--end", default=mt.default_end())
    args = ap.parse_args(argv)
    if args.vintages_per_batch < 1:
        ap.error("--vintages-per-batch muss >= 1 sein")

    today = date.today()
    full_plan = mt.ticker_plan()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    groups = batches(args.vintages_per_batch)

    print(f"{len(config.VINTAGES)} Jahrgaenge, {len(full_plan)} Ticker -> {len(groups)} Bloecke\n")
    for i, vs in enumerate(groups, start=1):
        plan = [r for r in full_plan if r["vintage"] in set(vs)]
        wb = Workbook()
        mt.build_daten_sheet(wb.active, plan, args.start, args.end)
        wb.active.title = f"Daten_Block{i}"
        mt.build_tickerliste_sheet(wb.create_sheet(), plan, today)
        mt.build_anleitung_sheet(wb.create_sheet())
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
        out = args.out_dir / f"pull_block{i}_{min(vs)}-{max(vs)}.xlsx"
        wb.save(out)
        print(f"Block {i}: Jahrgaenge {min(vs)}-{max(vs)}, {len(plan)} Ticker -> {out.name}")
        print(f"          {', '.join(r['ticker'] for r in plan)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
