"""
Synthetic Bloomberg exports for the RB/HO Jun-vs-Dec spread pipeline.

Generates realistic-looking PX_LAST histories for all 60 tickers in
``config.all_tickers()`` on the NYMEX trading calendar and writes

1. ``synthetic_export.xlsx``      Layout A "paired" (Excel BDH template): ticker in
                                  row 1, (date, value) column pairs below, prices in
                                  CENTS/gal, sheets "RB" and "HO" plus documentation
                                  sheets "Anleitung"/"Tickerliste" that must be skipped.
                                  Some tickers lower/upper case, some with a
                                  " COMB Comdty" suffix, some blocks with a
                                  "Date"/"PX_LAST" header row, some date columns as
                                  strings (dd.mm.yyyy / ISO) instead of date cells.
2. ``synthetic_export_wide.csv``  Layout B: first column date, two header rows
                                  (tickers, then "PX_LAST"), prices in $/gal.
3. ``synthetic_truth.csv``        the true spread_bbl per (date, vintage) computed from
                                  exactly the (rounded) prices written to the files, so
                                  tests can compare the pipeline output against it.

The 2026 Dec and 2027 Jun/Dec contracts use Bloomberg's ONE-digit ticker form
(XBZ6, HOZ6, XBM7, HOM7, XBZ7, HOZ7).

Price model ($/gal): HO = common mean-reverting level (shared by all contracts)
+ term-structure slope * months-to-expiry + contract-specific AR(1) noise.
RB = HO + crack; the December crack is a slow OU process per vintage, the June
crack = December crack + spread/42 where the spread ($/bbl) is an OU process
around a vintage-specific level (typically +5..+15) whose mean reversion
strengthens toward expiry (convergence). Vintages 2015 and 2020 are designed to
dip clearly negative for a stretch (2020 also expires negative), 2024 has a
short shallow dip. A few December quotes are removed (one of them on the June
expiry day of 2018) to exercise the inner join and the expiry logic.

Usage: py tests/make_synthetic.py [--seed N] [--out-dir DIR] [--as-of YYYY-MM-DD]
"""
from __future__ import annotations

import argparse
import sys
from bisect import bisect_right
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))
from spread import calendar_nymex, config  # noqa: E402

AS_OF_DEFAULT = date(2026, 9, 22)
DEFAULT_SEED = 20260922
XLSX_NAME = "synthetic_export.xlsx"
WIDE_CSV_NAME = "synthetic_export_wide.csv"
TRUTH_NAME = "synthetic_truth.csv"

LISTING_MONTHS_BEFORE_EXPIRY = 36
ONE_DIGIT_CONTRACTS = {("Z", 2026), ("M", 2027), ("Z", 2027)}   # active contracts shown one-digit by Bloomberg
COMB_SUFFIX_VINTAGES = {2014, 2019, 2024}                        # " COMB Comdty" instead of " Comdty"
# designed spread level profiles ($/bbl) as (tday_index within the window, level) knots
PROFILES = {
    2015: [(0, 6.0), (100, 6.0), (120, -5.0), (185, -5.0), (205, 3.0), (260, 4.0)],
    2020: [(0, 9.0), (180, 9.0), (195, -10.0), (230, -10.0), (245, -4.0), (260, -3.0)],
    2024: [(0, 8.0), (60, 8.0), (75, -1.5), (95, -1.5), (110, 7.0), (260, 9.0)],
}
# December quotes removed inside the window: (vintage, leg, "expiry" | "random")
MISSING_DEC = [(2018, "ho_dec", "expiry"), (2016, "rb_dec", "random"), (2019, "rb_dec", "random"),
               (2023, "rb_dec", "random"), (2021, "ho_dec", "random"), (2021, "ho_dec", "random")]
TRUTH_COLUMNS = ["date", "vintage", "rb_jun", "ho_jun", "rb_dec", "ho_dec", "spread_bbl"]


# ----------------------------------------------------------------------------
# stochastic building blocks
# ----------------------------------------------------------------------------
def _ou(rng, n, x0, mu, kappa, sigma, lo=None, hi=None) -> np.ndarray:
    x = np.empty(n)
    x[0] = x0
    z = rng.standard_normal(n)
    for t in range(1, n):
        x[t] = x[t - 1] + kappa * (mu - x[t - 1]) + sigma * z[t]
        if lo is not None and x[t] < lo:
            x[t] = lo + (lo - x[t])
        if hi is not None and x[t] > hi:
            x[t] = hi - (x[t] - hi)
    return x


def _ar1(rng, n, phi, sigma) -> np.ndarray:
    x = np.empty(n)
    x[0] = rng.normal(0, sigma / np.sqrt(1 - phi * phi))
    z = rng.standard_normal(n)
    for t in range(1, n):
        x[t] = phi * x[t - 1] + sigma * z[t]
    return x


def _profile(vintage: int, rng) -> list[tuple[int, float]]:
    if vintage in PROFILES:
        return PROFILES[vintage]
    mu0 = float(rng.uniform(5.0, 14.0))
    mid = mu0 + float(rng.normal(0.0, 2.0))
    end = float(rng.uniform(3.0, 15.0))
    return [(0, mu0), (150, mid), (260, end)]


def _spread_path(rng, vintage: int, grid: list[date], pos: dict[date, int], expiry_jun: date) -> np.ndarray:
    """Spread ($/bbl) of one vintage on the whole grid: OU around a piecewise-linear level,
    mean reversion strengthening toward the June expiry."""
    T = len(grid)
    w0 = pos[calendar_nymex.next_trading_day(config.window_start(vintage), include=True)]
    i_exp = pos[expiry_jun]
    knots = _profile(vintage, rng)
    kx = np.array([k[0] for k in knots], dtype=float)
    ky = np.array([k[1] for k in knots], dtype=float)
    j = np.arange(T) - w0
    level = np.interp(j, kx, ky)                      # constant extrapolation outside the knots
    tau = np.maximum(i_exp - np.arange(T), 0)
    kappa = 0.03 + 0.30 * np.exp(-tau / 12.0)
    sigma = 0.45 * (1.6 if vintage == 2020 else 1.0)
    x = np.empty(T)
    x[0] = level[0] + rng.normal(0, 1.5)
    z = rng.standard_normal(T)
    for t in range(1, T):
        x[t] = x[t - 1] + kappa[t] * (level[t] - x[t - 1]) + sigma * z[t]
    return x


# ----------------------------------------------------------------------------
# tickers
# ----------------------------------------------------------------------------
def raw_ticker(product: str, month_code: str, vintage: int, style: str) -> str:
    """Bloomberg ticker as it appears in the export ('xlsx' style adds case variations)."""
    root = config.PRODUCT_TO_BBG_ROOT[product]
    yy = f"{vintage % 10:d}" if (month_code, vintage) in ONE_DIGIT_CONTRACTS else f"{vintage % 100:02d}"
    suffix = " COMB Comdty" if vintage in COMB_SUFFIX_VINTAGES else " Comdty"
    t = f"{root}{month_code}{yy}{suffix}"
    if style == "xlsx":
        if vintage % 4 == 1:
            t = t.lower()
        elif vintage % 4 == 2 and product == "HO":
            t = t.replace("Comdty", "comdty")
        elif vintage % 4 == 3 and product == "RB":
            t = t.upper()
    return t


# ----------------------------------------------------------------------------
# generation
# ----------------------------------------------------------------------------
def _listing_start(rng, expiry: date, grid: list[date]) -> date:
    y, m = expiry.year, expiry.month - LISTING_MONTHS_BEFORE_EXPIRY
    while m <= 0:
        y, m = y - 1, m + 12
    d = date(y, m, 1) + timedelta(days=int(rng.integers(0, 15)))
    d = max(d, grid[0])
    return calendar_nymex.next_trading_day(d, include=True)


def generate(out_dir: Path | str, seed: int = DEFAULT_SEED, as_of: date = AS_OF_DEFAULT) -> dict:
    """Write the three files into out_dir and return their paths plus bookkeeping."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(seed)
    grid = calendar_nymex.trading_days(config.HISTORY_START, date(config.LAST_VINTAGE, 12, 31))
    pos = {d: i for i, d in enumerate(grid)}
    T = len(grid)
    n_obs = bisect_right(grid, as_of)                 # grid days <= as_of are observable
    level = _ou(rng, T, x0=2.9, mu=2.6, kappa=0.002, sigma=0.025, lo=1.4, hi=4.5)
    slope = _ou(rng, T, x0=0.0, mu=-0.003, kappa=0.01, sigma=0.0012, lo=-0.012, hi=0.012)

    series: dict[tuple[str, str, int], pd.Series] = {}   # (product, month_code, vintage) -> $/gal
    for v in config.VINTAGES:
        expiries = {mc: calendar_nymex.expected_expiry(v, mc) for mc in config.LEG_CODES}
        base_dec = float(rng.normal(-0.15, 0.06))
        dec_crack = _ou(rng, T, x0=base_dec, mu=base_dec, kappa=0.02, sigma=0.006)
        spread = _spread_path(rng, v, grid, pos, expiries[config.LEG_MONTHS["jun"]])
        crack = {"M": dec_crack + spread / config.GALLONS_PER_BBL, "Z": dec_crack}
        for mc in config.LEG_CODES:
            exp = expiries[mc]
            i0 = pos[_listing_start(rng, exp, grid)]
            i1 = min(pos[exp], n_obs - 1)
            if i0 > i1:
                continue
            sl = slice(i0, i1 + 1)
            months = np.array([(exp - d).days / 30.4375 for d in grid[sl]])
            ho = level[sl] + slope[sl] * months + _ar1(rng, i1 - i0 + 1, phi=0.98, sigma=0.006)
            rb = ho + crack[mc][sl]
            idx = pd.to_datetime([datetime(d.year, d.month, d.day) for d in grid[sl]]).astype("datetime64[ns]")
            for product, px in (("HO", ho), ("RB", rb)):
                cents = np.round(px * 100.0, 2)              # Bloomberg precision: 0.01 cent
                series[(product, mc, v)] = pd.Series(cents / 100.0, index=idx)

    # remove some December quotes inside the window
    missing: list[dict] = []
    for v, leg, how in MISSING_DEC:
        product, mc = next((p, m) for p, m, c in config.LEGS if c == leg)
        s = series[(product, mc, v)]
        if how == "expiry":
            d = pd.Timestamp(calendar_nymex.expected_expiry(v, config.LEG_MONTHS["jun"]))
        else:
            win = s.index[(s.index >= pd.Timestamp(config.window_start(v))) & (s.index <= pd.Timestamp(config.window_nominal_end(v)))]
            d = win[int(rng.integers(20, len(win) - 20))]
        s.loc[d] = np.nan
        missing.append({"vintage": v, "leg": leg, "ticker": config.bbg_ticker(product, mc, v), "date": d.date()})

    truth = _truth(series)
    paths = {
        "xlsx": out_dir / XLSX_NAME,
        "wide_csv": out_dir / WIDE_CSV_NAME,
        "truth": out_dir / TRUTH_NAME,
        "missing": missing,
        "as_of": as_of,
        "seed": seed,
        "n_tickers": len(series),
    }
    _write_xlsx(paths["xlsx"], series)
    _write_wide_csv(paths["wide_csv"], series)
    truth.to_csv(paths["truth"], index=False, date_format="%Y-%m-%d", lineterminator="\n")
    paths["truth_frame"] = truth
    return paths


def _truth(series: dict) -> pd.DataFrame:
    parts = []
    for v in config.VINTAGES:
        wide = pd.DataFrame({col: series[(p, mc, v)] for p, mc, col in config.LEGS if (p, mc, v) in series})
        if wide.shape[1] < len(config.LEGS):
            continue
        wide = wide[(wide.index >= pd.Timestamp(config.window_start(v))) & (wide.index <= pd.Timestamp(config.window_nominal_end(v)))]
        wide = wide.dropna().sort_index()
        g = config.GALLONS_PER_BBL
        wide["spread_bbl"] = ((wide["rb_jun"] - wide["ho_jun"]) - (wide["rb_dec"] - wide["ho_dec"])) * g
        wide["vintage"] = v
        wide.index.name = "date"
        parts.append(wide.reset_index())
    return pd.concat(parts, ignore_index=True)[TRUTH_COLUMNS]


# ----------------------------------------------------------------------------
# writers
# ----------------------------------------------------------------------------
def _write_xlsx(path: Path, series: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Anleitung"
    for line in ("Synthetischer Bloomberg-Export (Testdaten).",
                 "Blaetter RB und HO: BDH-Template, Ticker in Zeile 1, darunter Datum/PX_LAST Spaltenpaare.",
                 "Preise in US-Cent pro Gallone (Bloomberg-Konvention).",
                 "XBM13 Comdty", "Dieses Blatt wird vom Loader uebersprungen."):
        ws.append([line])
    for product in config.PRODUCTS:
        ws = wb.create_sheet(product)
        keys = [(p, mc, v) for v in config.VINTAGES for p, mc, _ in config.LEGS if p == product and (p, mc, v) in series]
        for i, key in enumerate(keys):
            _, mc, v = key
            col = 2 * i + 1
            ws.cell(row=1, column=col, value=raw_ticker(product, mc, v, "xlsx"))
            r0 = 2
            if i % 3 == 0:                          # Bloomberg-style header row in some blocks
                ws.cell(row=2, column=col, value="Date")
                ws.cell(row=2, column=col + 1, value="PX_LAST")
                r0 = 3
            style = i % 7
            s = series[key]
            for j, (ts, val) in enumerate(zip(s.index, s.to_numpy())):
                r = r0 + j
                if style == 3:
                    cell = ts.strftime("%d.%m.%Y")
                elif style == 5:
                    cell = ts.strftime("%Y-%m-%d")
                else:
                    cell = ts.to_pydatetime()
                ws.cell(row=r, column=col, value=cell)
                if not np.isnan(val):
                    ws.cell(row=r, column=col + 1, value=round(float(val) * 100.0, 2))
    ws = wb.create_sheet("Tickerliste")
    ws.append(["Ticker", "Beschreibung"])
    for (p, mc, v) in sorted(series):
        ws.append([raw_ticker(p, mc, v, "csv"), f"{p} {mc}{v}"])
    wb.save(path)


def _write_wide_csv(path: Path, series: dict) -> None:
    order = [(p, mc, v) for v in config.VINTAGES for p, mc, _ in config.LEGS if (p, mc, v) in series]
    wide = pd.DataFrame({raw_ticker(p, mc, v, "csv"): series[(p, mc, v)] for p, mc, v in order}).sort_index()
    wide.index = [ts.strftime("%Y-%m-%d") for ts in wide.index]
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("Date," + ",".join(wide.columns) + "\n")
        fh.write("," + ",".join(["PX_LAST"] * wide.shape[1]) + "\n")
        wide.to_csv(fh, header=False, index=True, float_format="%.4f", na_rep="", lineterminator="\n")


# ----------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--seed", type=int, default=DEFAULT_SEED)
    ap.add_argument("--out-dir", type=Path, default=config.DATA_RAW)
    ap.add_argument("--as-of", type=date.fromisoformat, default=AS_OF_DEFAULT, help="last observable date (default 2026-09-22)")
    a = ap.parse_args(argv)
    res = generate(a.out_dir, seed=a.seed, as_of=a.as_of)
    truth = res["truth_frame"]
    neg = truth.groupby("vintage")["spread_bbl"].agg(lambda s: int((s < 0).sum()))
    last = truth.groupby("vintage")["spread_bbl"].last()
    print(f"seed {a.seed}, as-of {a.as_of}, {res['n_tickers']} tickers")
    for k in ("xlsx", "wide_csv", "truth"):
        print(f"  {res[k]}")
    print(f"truth: {len(truth)} rows, {truth['vintage'].nunique()} vintages, {truth['date'].min().date()} .. {truth['date'].max().date()}")
    print("negative days per vintage: " + ", ".join(f"{v}:{n}" for v, n in neg.items() if n))
    print("last value per vintage:    " + ", ".join(f"{v}:{x:+.1f}" for v, x in last.items()))
    print("removed December quotes:   " + ", ".join(f"{m['ticker']} {m['date']}" for m in res["missing"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
